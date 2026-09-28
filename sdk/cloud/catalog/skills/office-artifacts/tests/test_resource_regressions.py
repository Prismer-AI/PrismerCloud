"""Offline regression gates for the resource audit (unittest, no account access)."""
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1] / 'references'


def module(kind, name):
    path = ROOT / kind / 'scripts' / (name + '.py')
    sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


class ResourceRegression(unittest.TestCase):
    def test_pdf_rasterizer_runtime_fallback_and_boundaries(self):
        raster = module('pdf', '_raster')
        with patch.object(raster, 'available_backends', return_value=['pypdfium2', 'pdftoppm']), patch.object(raster, '_via_pdfium', side_effect=RuntimeError('renderer unavailable')), patch.object(raster, '_via_pdftoppm', return_value='image') as fallback:
            self.assertEqual(raster.rasterize_page('fake.pdf', 1), 'image')
            fallback.assert_called_once()
        with self.assertRaises(ValueError): raster.rasterize_page('fake.pdf', 0)
        with self.assertRaises(ValueError): raster.rasterize_page('fake.pdf', 1, dpi=100000)

    def test_docx_existing_generic_comment_part_persists_deletion(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        comments = module('docx', 'docx_comments')
        d = Document(); d.add_paragraph('text')
        blob = f'<w:comments xmlns:w="{comments.W}"><w:comment w:id="0"><w:p><w:r><w:t>old</w:t></w:r></w:p></w:comment></w:comments>'.encode()
        part = Part(PackURI('/word/comments.xml'), comments.COMMENTS_CT, blob, d.part.package)
        d.part.relate_to(part, RT.COMMENTS)
        self.assertTrue(comments.delete_comment(d, '0'))
        path = self.dir / 'comments.docx'; d.save(path)
        self.assertEqual(comments.list_comments(Document(path)), [])

    def test_pdf_unicode_requires_explicit_font_and_empty_split_rejected(self):
        create = module('pdf', 'pdf_create')
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertNotEqual(create.build_pdf({'elements': [{'type': 'paragraph', 'text': '\u4e2d\u6587'}]}, str(self.dir / 'bad.pdf')), 0)
        self.assertFalse((self.dir / 'bad.pdf').exists())
        split = module('pdf', 'pdf_split')
        with self.assertRaises(ValueError): split.parse_pages(' , ', 2)

    def test_xlsx_table_partial_column_change_fails_before_output(self):
        w = Workbook(); s = w.active
        s.append(['first', 'second']); s.append([1, 2]); s.add_table(Table(displayName='T', ref='A1:B2'))
        src = self.dir / 'table.xlsx'; dest = self.dir / 'edited.xlsx'; w.save(src)
        result = self.cli('xlsx', 'xlsx_restructure', src, '--delete-cols', 'A', '--out', dest)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(dest.exists())
        self.assertIn('table', result.stderr.lower())

    def test_xlsx_insert_table_column_keeps_schema_and_filter_consistent(self):
        w = Workbook(); s = w.active
        s.append(['first', 'second']); s.append([1, 2]); s.add_table(Table(displayName='T', ref='A1:B2'))
        src = self.dir / 'table.xlsx'; dest = self.dir / 'edited.xlsx'; w.save(src)
        result = self.cli('xlsx', 'xlsx_restructure', src, '--insert-cols', 'B', '--out', dest)
        self.assertEqual(result.returncode, 0, result.stderr)
        s = load_workbook(dest).active; t = s.tables['T']
        self.assertEqual(t.ref, 'A1:C2'); self.assertEqual(t.autoFilter.ref, t.ref)
        self.assertEqual([c.name for c in t.tableColumns], [s.cell(1, i).value for i in range(1, 4)])
        self.assertEqual(len(t.tableColumns), 3)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)

    def cli(self, kind, name, *args):
        return subprocess.run([sys.executable, '-B', str(ROOT / kind / 'scripts' / (name + '.py')),
                               *map(str, args)], capture_output=True, text=True, timeout=20)

    def test_docx_replacement_terminates_without_rescanning_new_text(self):
        code = "from docx import Document; from docx_common import replace_in_paragraph; p=Document().add_paragraph('x x'); assert replace_in_paragraph(p,'x','xx')==2; assert p.text=='xx xx'"
        try:
            result = subprocess.run([sys.executable, '-B', '-c', code], timeout=2,
                                    env={**os.environ, 'PYTHONPATH': str(ROOT / 'docx/scripts')}, capture_output=True)
        except subprocess.TimeoutExpired:
            self.fail('replacement rescans generated text and does not terminate')
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_docx_accepted_text_and_namespace_independent_revision_detection(self):
        d = Document(); p = d.add_paragraph('before ')
        ins = OxmlElement('w:ins'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
        t.text = 'inserted'; r.append(t); ins.append(r); p._p.append(ins)
        reader = module('docx', 'docx_read')
        self.assertEqual(reader.extract_text(d)['body'], ['before inserted'])
        p._p.remove(ins); p._p.append(OxmlElement('w:moveFrom'))
        dest = self.dir / 'move.docx'; d.save(dest)
        self.assertTrue(reader.detect_revisions(dest)['has_tracked_changes'])

    def test_docx_structural_revision_is_not_silently_resolved(self):
        d = Document(); p = d.add_paragraph('text')
        props = p._p.get_or_add_pPr(); props.append(OxmlElement('w:del'))
        revision = module('docx', 'docx_revisions')
        self.assertEqual(revision.resolve(d, True), 0)
        self.assertIsNotNone(props.find(revision.q('del')))

    def test_docx_validator_keeps_relationship_source_parts_separate(self):
        image = self.dir / 'tiny.png'; Image.new('RGB', (3, 3), 'red').save(image)
        d = Document(); d.add_picture(str(image)); d.sections[0].header.paragraphs[0].add_run().add_picture(str(image))
        path = self.dir / 'doc.docx'; d.save(path)
        for seed in range(5):
            code = f"from docx_validate import validate; import json; print(json.dumps(validate({str(path)!r})))"
            p = subprocess.run([sys.executable, '-B', '-c', code], capture_output=True, text=True,
                               env={**os.environ, 'PYTHONPATH': str(ROOT / 'docx/scripts'), 'PYTHONHASHSEED': str(seed)})
            self.assertTrue(json.loads(p.stdout)['ok'], p.stdout)

    def test_pdf_flatten_removes_fields_and_widgets_but_retains_value(self):
        src = self.dir / 'form.pdf'; out = self.dir / 'flat.pdf'; fields = self.dir / 'values.json'
        c = Canvas(str(src)); c.acroForm.textfield(name='name', x=20, y=700, width=100, height=20); c.showPage(); c.save()
        fields.write_text('{"name":"Alice"}')
        p = self.cli('pdf', 'pdf_fill_form', src, '--fields-json', fields, '-o', out, '--flatten')
        self.assertEqual(p.returncode, 0, p.stderr)
        doc = PdfReader(out)
        self.assertFalse(doc.get_fields())
        self.assertFalse(doc.pages[0].get('/Annots'))
        self.assertIn('Alice', doc.pages[0].extract_text())

    def test_pdf_attachment_names_cannot_overwrite_existing_or_drop_duplicates(self):
        src = self.dir / 'attachments.pdf'; dest = self.dir / 'out'; dest.mkdir()
        (dest / 'same.txt').write_bytes(b'existing')
        writer = PdfWriter(); writer.add_blank_page(100, 100)
        writer.add_attachment('same.txt', b'first'); writer.add_attachment('same.txt', b'second'); writer.write(src)
        result = self.cli('pdf', 'pdf_meta', src, '--extract-attachments', dest)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual((dest / 'same.txt').read_bytes(), b'existing')
        self.assertEqual({p.read_bytes() for p in dest.iterdir()}, {b'existing', b'first', b'second'})

    def test_csv_literal_mode_does_not_create_formulas(self):
        source = self.dir / 'in.csv'; out = self.dir / 'out.xlsx'; source.write_text('value\n=1+1\n')
        self.assertEqual(self.cli('xlsx', 'csv_to_xlsx', source, out, '--no-infer').returncode, 0)
        self.assertEqual(load_workbook(out).active['A2'].data_type, 's')

    def test_xlsx_formula_token_types_and_whole_column_ranges(self):
        rw = module('xlsx', 'xlsx_restructure').RefRewriter('Sheet1', 'cols', 1, 1, False)
        for source, expected in [('=SUM(B:B)', '=SUM(C:C)'), ('=SUM(Table1[A1])', '=SUM(Table1[A1])'),
                                 ('="A1"&A1', '="A1"&B1'), ('=SUM(1:2)', '=SUM(1:2)'),
                                 ("='[other.xlsx]Sheet1'!A1", "='[other.xlsx]Sheet1'!A1")]:
            self.assertEqual(rw.rewrite(source, 'Sheet1'), expected)

    def test_xlsx_column_dimensions_survive_shift(self):
        src = self.dir / 'in.xlsx'; w = Workbook(); w.active.column_dimensions['B'].width = 24; w.save(src)
        p = self.cli('xlsx', 'xlsx_restructure', src, '--insert-cols', 'B')
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertEqual(load_workbook(src).active.column_dimensions['C'].width, 24)

    def test_xlsx_full_delete_removes_validation_table_and_dimension(self):
        src = self.dir / 'in.xlsx'; w = Workbook(); s = w.active
        s.append(['name']); s.append(['value']); s.add_table(Table(displayName='Items', ref='A1:A2'))
        dv = DataValidation(type='whole'); dv.add('A1:A2'); s.add_data_validation(dv); s.row_dimensions[1].height = 30; w.save(src)
        p = self.cli('xlsx', 'xlsx_restructure', src, '--delete-rows', '1:2')
        self.assertEqual(p.returncode, 0, p.stderr); s = load_workbook(src).active
        self.assertFalse(s.tables); self.assertFalse(s.data_validations.dataValidation); self.assertNotIn(1, s.row_dimensions)

    def test_pptx_mixed_runs_and_group_text_are_all_replaced(self):
        prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6])
        tf = s.shapes.add_group_shape().shapes.add_textbox(0, 0, 1000000, 1000000).text_frame
        for value in ['foo ', 'f', 'oo']: tf.paragraphs[0].add_run().text = value
        self.assertEqual(module('powerpoint', 'pptx_edit').replace_text(prs, 'foo', 'bar'), 2)
        self.assertEqual(tf.text, 'bar bar')

    def test_pptx_copy_preserves_background(self):
        prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6]); edit = module('powerpoint', 'pptx_edit')
        edit.set_background(s, 'FF0000'); edit.duplicate_slide(prs, 0)
        self.assertEqual(prs.slides[1].background.fill.type, s.background.fill.type)

    def test_pptx_grouped_chart_copy_is_refused(self):
        prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6]); data = CategoryChartData()
        data.categories = ['A']; data.add_series('S', [1])
        chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 1000000, 1000000, data)
        s.shapes.add_group_shape([chart])
        with self.assertRaises(SystemExit): module('powerpoint', 'pptx_edit').duplicate_slide(prs, 0)

    def test_render_returns_only_current_generation(self):
        out = self.dir / 'render'; out.mkdir(); (out / 'slide-99.png').write_bytes(b'stale')
        renderer = module('powerpoint', 'pptx_render')
        def fake_run(cmd, **kwargs):
            if '--outdir' in cmd: (Path(cmd[cmd.index('--outdir') + 1]) / 'deck.pdf').write_bytes(b'pdf')
            else: Path(cmd[-1] + '-1.png').write_bytes(b'current')
            return subprocess.CompletedProcess(cmd, 0, '', '')
        with patch.object(renderer, 'find_tools', return_value=('soffice', 'pdftoppm', [])), patch.object(renderer.subprocess, 'run', side_effect=fake_run):
            result = renderer.render('deck.pptx', str(out), 'slide', 100)
        self.assertEqual(len(result['files']), 1)
        self.assertEqual(Path(result['files'][0]).read_bytes(), b'current')

    def test_pdf_reverse_and_empty_ranges_are_errors(self):
        for name in ('pdf_page_image', 'pdf_stamp'):
            helper = module('pdf', name)
            for value in ('3-1', '', '0-0'):
                with self.assertRaises(ValueError): helper.parse_pages(value, 3)

    def test_docx_strict_template_failure_does_not_publish(self):
        src = self.dir / 'in.docx'; out = self.dir / 'out.docx'; values = self.dir / 'values.json'
        d = Document(); d.add_paragraph('{{missing}}'); d.save(src); values.write_text('{}')
        self.assertNotEqual(self.cli('docx', 'docx_template', src, values, out, '--strict').returncode, 0)
        self.assertFalse(out.exists())

    def test_docx_linked_headers_are_not_modified_twice(self):
        d = Document(); d.sections[0].header.paragraphs[0].text = 'x'; d.add_section()
        common = module('docx', 'docx_common')
        for paragraph in common.iter_all_paragraphs(d): common.replace_in_paragraph(paragraph, 'x', 'xx')
        self.assertEqual(d.sections[0].header.paragraphs[0].text, 'xx')

    def test_table_append_cannot_overwrite_cells_or_leave_stale_filter(self):
        w = Workbook(); s = w.active; s.append(['name']); s.append(['first'])
        table = Table(displayName='Items', ref='A1:A2'); s.add_table(table)
        src = self.dir / 'table.xlsx'; w.save(src); s = load_workbook(src).active
        edit = module('xlsx', 'xlsx_edit'); s['A3'] = 'existing'
        with self.assertRaises(ValueError): edit.table_append(s, 'Items', ['second'])
        s['A3'] = None; edit.table_append(s, 'Items', ['second'])
        self.assertEqual(s.tables['Items'].autoFilter.ref, 'A1:A3')

    def test_pymupdf_page_selection_applies_to_tables(self):
        import types
        helper = module('pdf', 'extract_pymupdf'); visited = []
        class Page:
            def __init__(self, number): self.number = number
            def find_tables(self): visited.append(self.number); return types.SimpleNamespace(tables=[])
        with patch.dict(sys.modules, {'pymupdf': types.SimpleNamespace(open=lambda p: [Page(0), Page(1)])}):
            helper.extract_tables('fake.pdf', pages=[1])
        self.assertEqual(visited, [1])

    def test_marker_json_renderer_and_pil_image_output(self):
        import types
        helper = module('pdf', 'extract_marker'); options = []
        class Config:
            def __init__(self, config): options.append(config)
            def generate_config_dict(self): return options[-1]
            def get_renderer(self): return 'renderer'
            def get_processors(self): return []
            def get_llm_service(self): return None
        rendered = types.SimpleNamespace(model_dump=lambda: {'children': [{'text': 'hello'}]})
        def converter(**kwargs):
            self.assertEqual(kwargs.get('renderer'), 'renderer')
            return lambda path: rendered
        modules = {'marker.converters.pdf': types.SimpleNamespace(PdfConverter=converter),
                   'marker.models': types.SimpleNamespace(create_model_dict=lambda: {}),
                   'marker.config.parser': types.SimpleNamespace(ConfigParser=Config),
                   'marker.output': types.SimpleNamespace(text_from_rendered=lambda obj: ('', 'json', {'image.png': Image.new('RGB', (3, 3))}))}
        output = io.StringIO()
        with patch.dict(sys.modules, modules), contextlib.redirect_stdout(output): helper.convert('fake.pdf', str(self.dir), 'json')
        self.assertEqual(json.loads(output.getvalue())['children'][0]['text'], 'hello')
        self.assertTrue((self.dir / 'image.png').exists())


if __name__ == '__main__':
    unittest.main()
